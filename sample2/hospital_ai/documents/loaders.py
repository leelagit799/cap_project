"""Raw content extraction from clinical source documents.

Doc §2.2 requires multi-modal, multi-language extraction; Table 14 lists
".txt / .pdf / .docx" plus optional Tesseract OCR. The shipped sample data adds
JSON exports and scanned PNGs.

Scanned documents ship with pre-generated ``<name>.ocr.txt`` sidecars. Those are
preferred over running Tesseract: they are the transcription the dataset was
built against, so extraction stays deterministic whether or not the OCR binary
is installed.
"""

from __future__ import annotations

import hashlib
import json
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from hospital_ai.core.errors import ExtractionError
from hospital_ai.core.logging import get_logger

_log = get_logger(__name__, component="document-loader")

TEXT_SUFFIXES = {".txt", ".text", ".md"}
JSON_SUFFIXES = {".json"}
PDF_SUFFIXES = {".pdf"}
DOCX_SUFFIXES = {".docx"}
IMAGE_SUFFIXES = {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp"}

SUPPORTED_SUFFIXES = (
    TEXT_SUFFIXES | JSON_SUFFIXES | PDF_SUFFIXES | DOCX_SUFFIXES | IMAGE_SUFFIXES
)

#: Several patients ship the same document in more than one encoding — P1020's
#: bill exists as both a PDF and a JSON export. Lower sorts first, so a
#: machine-readable form is preferred over one that needs OCR.
ENCODING_PRIORITY = {".json": 0, ".txt": 1, ".docx": 2, ".pdf": 3, ".png": 4, ".jpg": 5}


def encoding_rank(path: Path) -> int:
    return ENCODING_PRIORITY.get(path.suffix.lower(), 9)


MEDIA_TYPES = {
    ".txt": "text/plain",
    ".text": "text/plain",
    ".md": "text/markdown",
    ".json": "application/json",
    ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".png": "image/png",
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
    ".tif": "image/tiff",
    ".tiff": "image/tiff",
    ".bmp": "image/bmp",
}


@dataclass
class DocumentContent:
    """Everything the Harvester Tool extracts from one file."""

    text: str
    media_type: str
    tables: list[list[list[str]]] = field(default_factory=list)
    structured: dict[str, Any] | None = None
    page_count: int = 1
    ocr_used: bool = False
    ocr_source: str | None = None
    warnings: list[str] = field(default_factory=list)

    @property
    def is_empty(self) -> bool:
        return not self.text.strip() and self.structured is None


def sha256_of(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(65536), b""):
            digest.update(block)
    return digest.hexdigest()


def media_type_for(path: Path) -> str:
    return MEDIA_TYPES.get(path.suffix.lower(), "application/octet-stream")


def is_ocr_sidecar(path: Path) -> bool:
    """True for ``P1020_labs.pdf.ocr.txt`` — a transcription, not a source doc."""
    return path.name.endswith(".ocr.txt")


def find_ocr_sidecar(path: Path) -> Path | None:
    sidecar = path.with_name(path.name + ".ocr.txt")
    return sidecar if sidecar.is_file() else None


# --- Per-format loaders ------------------------------------------------------


def _load_text(path: Path) -> DocumentContent:
    return DocumentContent(
        text=path.read_text(encoding="utf-8", errors="replace"),
        media_type=media_type_for(path),
    )


def _load_json(path: Path) -> DocumentContent:
    try:
        payload = json.loads(path.read_text(encoding="utf-8"))
    except json.JSONDecodeError as exc:
        raise ExtractionError(f"{path.name}: invalid JSON ({exc.msg} at line {exc.lineno})") from exc

    # Several sample bills and lab reports carry a `raw_text` narrative
    # alongside the structured fields; it is the better basis for translation
    # and embedding than a re-serialised dict.
    narrative = payload.get("raw_text") if isinstance(payload, dict) else None
    text = narrative if isinstance(narrative, str) and narrative.strip() else json.dumps(
        payload, indent=2, ensure_ascii=False
    )
    return DocumentContent(
        text=text,
        media_type="application/json",
        structured=payload if isinstance(payload, dict) else {"items": payload},
    )


def _load_pdf(path: Path) -> DocumentContent:
    try:
        import pdfplumber
    except ImportError as exc:  # pragma: no cover - dependency is pinned
        raise ExtractionError("pdfplumber is required to read PDF documents") from exc

    pages: list[str] = []
    tables: list[list[list[str]]] = []
    warnings: list[str] = []

    with pdfplumber.open(path) as pdf:
        page_count = len(pdf.pages)
        for page in pdf.pages:
            pages.append(page.extract_text() or "")
            for table in page.extract_tables() or []:
                tables.append([[cell or "" for cell in row] for row in table])

    text = "\n".join(pages).strip()
    content = DocumentContent(
        text=text,
        media_type="application/pdf",
        tables=tables,
        page_count=page_count,
        warnings=warnings,
    )

    if not text:
        # A scanned PDF with no text layer.
        ocr_text, source = _ocr_fallback(path)
        if ocr_text:
            content.text = ocr_text
            content.ocr_used = True
            content.ocr_source = source
        else:
            content.warnings.append(
                f"{path.name}: no text layer and no OCR transcription available"
            )
    return content


def _load_docx(path: Path) -> DocumentContent:
    try:
        import docx
    except ImportError as exc:  # pragma: no cover - dependency is pinned
        raise ExtractionError("python-docx is required to read DOCX documents") from exc

    document = docx.Document(str(path))
    paragraphs = [p.text for p in document.paragraphs if p.text.strip()]
    tables = [
        [[cell.text.strip() for cell in row.cells] for row in table.rows]
        for table in document.tables
    ]
    return DocumentContent(
        text="\n".join(paragraphs),
        media_type=media_type_for(path),
        tables=tables,
    )


def _load_image(path: Path) -> DocumentContent:
    text, source = _ocr_fallback(path)
    content = DocumentContent(
        text=text or "",
        media_type=media_type_for(path),
        ocr_used=bool(text),
        ocr_source=source,
    )
    if not text:
        content.warnings.append(
            f"{path.name}: no OCR sidecar and Tesseract produced no text"
        )
    return content


def _ocr_fallback(path: Path) -> tuple[str, str | None]:
    """Prefer the shipped transcription; fall back to live Tesseract."""
    sidecar = find_ocr_sidecar(path)
    if sidecar is not None:
        return sidecar.read_text(encoding="utf-8", errors="replace"), f"sidecar:{sidecar.name}"

    try:
        import pytesseract
        from PIL import Image
    except ImportError:
        return "", None

    try:
        if path.suffix.lower() in PDF_SUFFIXES:
            # Rasterising PDFs needs poppler, which the spec does not require.
            return "", None
        with Image.open(path) as image:
            return pytesseract.image_to_string(image), "tesseract"
    except Exception as exc:  # noqa: BLE001 - OCR is best-effort by design
        _log.warning("OCR failed", extra={"document": path.name, "error": str(exc)})
        return "", None


_LOADERS = [
    (TEXT_SUFFIXES, _load_text),
    (JSON_SUFFIXES, _load_json),
    (PDF_SUFFIXES, _load_pdf),
    (DOCX_SUFFIXES, _load_docx),
    (IMAGE_SUFFIXES, _load_image),
]


def load_document(path: Path) -> DocumentContent:
    """Extract text, tables and structured data from one clinical document."""
    if not path.is_file():
        raise ExtractionError(f"Document not found: {path}")

    suffix = path.suffix.lower()
    for suffixes, loader in _LOADERS:
        if suffix in suffixes:
            content = loader(path)
            _log.debug(
                "document loaded",
                extra={
                    "document": path.name,
                    "media_type": content.media_type,
                    "chars": len(content.text),
                    "tables": len(content.tables),
                    "ocr_used": content.ocr_used,
                },
            )
            return content

    raise ExtractionError(
        f"{path.name}: unsupported format '{suffix}'. "
        f"Supported: {', '.join(sorted(SUPPORTED_SUFFIXES))}"
    )
